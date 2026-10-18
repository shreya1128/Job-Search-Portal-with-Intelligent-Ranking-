#!/usr/bin/env python3
"""Agreement between two labelled workbooks (e.g. HUMAN vs AI-judged) on the jobs both have labelled.

    python scripts/compare_labels.py --a my_labels.xlsx --b relevance_judgments_assistant_labeled.xlsx [--name-a human --name-b assistant_judged]

Rows are matched on (query_id, source, external_id); only rows labelled in BOTH files count.  Prints the report and writes label_agreement.json next to the first file."""
import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook

from jobintel.eval.agreement import agreement_report


def read(path):
    rows = list(load_workbook(path, read_only=True)["Judgments"].iter_rows(values_only=True))
    h = list(rows[0])
    q, s, e, r = (h.index(c) for c in ("query_id", "source", "external_id", "relevance"))
    out = {}
    for x in rows[1:]:
        if x[r] in (0, 1, 2, "0", "1", "2"):
            out[(str(x[q]), str(x[s]), str(x[e]))] = int(x[r])
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--a", required=True)
    ap.add_argument("--b", required=True)
    ap.add_argument("--name-a", default="A")
    ap.add_argument("--name-b", default="B")
    a = ap.parse_args()
    A, B = read(a.a), read(a.b)
    both = sorted(set(A) & set(B))
    print(f"{a.name_a}: {len(A)} labelled rows | {a.name_b}: {len(B)} | labelled in both: {len(both)}")
    if not both:
        sys.exit("no row is labelled in both files")
    rep = agreement_report([(k[0], A[k], B[k]) for k in both])
    print(f"exact agreement {rep['exact_agreement']:.3f} | within one grade {rep['within_one_grade']:.3f} | kappa {rep['kappa']:.3f} | weighted kappa (quadratic) {rep['weighted_kappa_quadratic']:.3f}"
          + (f"  95% CI over queries [{rep['weighted_kappa_ci95_over_queries'][0]:.3f}, {rep['weighted_kappa_ci95_over_queries'][1]:.3f}]" if 'weighted_kappa_ci95_over_queries' in rep else ""))
    print(f"mean grade: {a.name_a} {rep['mean_label']['a']:.2f}  vs  {a.name_b} {rep['mean_label']['b']:.2f}")
    for t in (1, 2):
        b = rep[f"binary_thr{t}"]
        print(f"relevant = grade >= {t}: agreement {b['agreement']:.3f}, kappa {b['kappa']:.3f}, share relevant {b['share_relevant_a']:.2f} vs {b['share_relevant_b']:.2f}")
    print("confusion (rows = " + a.name_a + ", columns = " + a.name_b + "):", rep["confusion_rows_a_cols_b"])
    for q, v in rep["per_query"].items():
        print(f"  {q}: n={v['n']} exact={v['exact']:.2f} mean {v['mean_a']:.2f} vs {v['mean_b']:.2f}")
    (Path(a.a).parent / "label_agreement.json").write_text(json.dumps(rep, indent=1))
