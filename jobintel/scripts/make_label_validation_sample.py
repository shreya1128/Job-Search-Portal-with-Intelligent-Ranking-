#!/usr/bin/env python3
"""Make a SMALL BLANK sample of the relevance workbook for a human to label, so that the AI-judged labels can be validated (scripts/compare_labels.py).

    python scripts/make_label_validation_sample.py --source relevance_judgments_assistant_labeled.xlsx --out relevance_validation_sample_TO_LABEL.xlsx

Default queries q06 q08 q11 q16 q19 (~115 rows, roughly 30 minutes) were chosen because the AI labels contain the most mixed grades there (the hardest judgments), so the agreement estimate
is conservative rather than flattering.  The AI's labels are NOT copied into the sample (no anchoring).  Same columns and instructions as the full workbook; fill ONLY `relevance` (0 / 1 / 2)."""
import argparse
from pathlib import Path

from openpyxl import load_workbook

NOTE = ("VALIDATION SAMPLE - label ONLY the `relevance` column on the Judgments sheet (0 = not relevant, 1 = partly, 2 = clearly relevant), using the rules below, judging each row on its own "
        "against the query's information need.  About 30 minutes.  Return the file; it is compared with the AI-judged labels (Cohen's kappa) - it does not change any ranking result by itself.")

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--queries", nargs="+", default=["q06", "q08", "q11", "q16", "q19"])
    a = ap.parse_args()
    wb = load_workbook(a.source)
    ws = wb["Judgments"]
    h = [c.value for c in ws[1]]
    qi, ri, ni = h.index("query_id") + 1, h.index("relevance") + 1, h.index("notes") + 1
    keep = set(a.queries)
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, qi).value not in keep:
            ws.delete_rows(r)
        else:
            ws.cell(r, ri).value = None
            ws.cell(r, ni).value = None
    qs = wb["Queries"]
    for r in range(qs.max_row, 1, -1):
        if qs.cell(r, 1).value not in keep:
            qs.delete_rows(r)
    ins = wb["Instructions"]
    ins.insert_rows(1)
    ins.cell(1, 1).value = NOTE
    wb.save(a.out)
    print(f"wrote {a.out}: {ws.max_row - 1} rows to label for queries {sorted(keep)}")
