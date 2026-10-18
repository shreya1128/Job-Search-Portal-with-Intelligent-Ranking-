import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from jobintel.eval import judgments as J
from jobintel.eval import runner as R
from jobintel.eval.report import render
from tests.test_ir_engine import JOBS, make_engine

QUERIES = [{"id": "q1", "text": "python sql", "need": "python and sql roles"}, {"id": "q2", "text": "nurse patient care", "need": "nurse roles"}]


def systems_and_engine():
    from jobintel.eval.systems import build_systems
    eng = make_engine(pagerank_values=[0.1, 0.2, 0.3, 0.15, 0.25])
    return eng, build_systems(eng.index, eng.docstore, eng.extractor, eng.lex, k=5)


def test_pool_is_shuffled_deduplicated_and_hides_the_system(tmp_path):
    eng, sy = systems_and_engine()
    pool = J.build_pool(sy, QUERIES, ["keyword_baseline", "lm_dirichlet", "dfr_pl2", "hybrid"], depth=3)
    for qid, p in pool.items():
        assert len(p["job_ids"]) == len(set(p["job_ids"])) and set(p["job_ids"]) == {i for v in p["systems"].values() for i in v}
    assert J.build_pool(sy, QUERIES, ["lm_dirichlet", "dfr_pl2"], 3) == J.build_pool(sy, QUERIES, ["lm_dirichlet", "dfr_pl2"], 3)           # reproducible
    rows = J.sheet_rows(pool, eng.index, eng.docstore)
    J.write_sheet(rows, QUERIES, tmp_path / "s.xlsx")
    wb = load_workbook(tmp_path / "s.xlsx")
    head = [c.value for c in wb["Judgments"][1]]
    assert head == J.HEADER and not any(h in head for h in ("system", "rank", "score", "model"))        # nothing that reveals which method found a job
    assert wb["Judgments"].data_validations.dataValidation and wb["Instructions"]["A1"].value.startswith("HOW TO LABEL")


def fill(path, grades):
    wb = load_workbook(path)
    ws = wb["Judgments"]
    col = J.HEADER.index("relevance") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(r, col).value = grades(ws.cell(r, 1).value, int(ws.cell(r, 4).value))
    wb.save(path)


def test_labels_roundtrip_and_validation(tmp_path):
    eng, sy = systems_and_engine()
    pool = J.build_pool(sy, QUERIES, ["lm_dirichlet", "dfr_pl2"], 4)
    J.write_sheet(J.sheet_rows(pool, eng.index, eng.docstore), QUERIES, tmp_path / "s.xlsx")
    labels, st0 = J.read_labels(tmp_path / "s.xlsx")
    assert st0["judged"] == 0 and st0["unlabeled"] == sum(len(p["job_ids"]) for p in pool.values())                      # a blank sheet is detected as unlabeled
    fill(tmp_path / "s.xlsx", lambda q, j: 2 if "Machine" in eng.index.meta["title"][eng.index.doc_index(j)] else 0)
    labels, st = J.read_labels(tmp_path / "s.xlsx")
    assert st["unlabeled"] == 0 and st["judged"] == st0["unlabeled"] and set(st["grade_counts"]) == {0, 1, 2}
    fill(tmp_path / "s.xlsx", lambda q, j: 7)
    with pytest.raises(ValueError, match="must be 0, 1 or 2"):
        J.read_labels(tmp_path / "s.xlsx")
    J.write_csv(J.sheet_rows(pool, eng.index, eng.docstore), tmp_path / "s.csv")
    assert J.read_labels(tmp_path / "s.csv")[1]["unlabeled"] == st0["unlabeled"]


def test_end_to_end_evaluation_with_known_labels(tmp_path):
    eng, sy = systems_and_engine()
    pool = J.build_pool(sy, QUERIES, ["keyword_baseline", "lm_dirichlet", "dfr_pl2", "hybrid", "hybrid_no_pagerank"], 4)
    J.write_sheet(J.sheet_rows(pool, eng.index, eng.docstore), QUERIES, tmp_path / "s.xlsx")
    title = lambda j: eng.index.meta["title"][eng.index.doc_index(j)]
    fill(tmp_path / "s.xlsx", lambda q, j: (2 if q == "q2" and "Nurse" in title(j) else 2 if q == "q1" and "Nurse" not in title(j) else 0))
    labels, st = J.read_labels(tmp_path / "s.xlsx")
    labels, missing = J.resolve_labels(labels, eng.index)
    assert missing == 0
    df = R.evaluate(sy, ["keyword_baseline", "lm_dirichlet", "dfr_pl2", "hybrid", "hybrid_no_pagerank"], QUERIES, labels)
    assert set(df.columns) >= set(R.METRICS) and len(df) == 10 and df["judged@10"].min() > 0
    nurse = df[(df.query_id == "q2") & (df.system == "lm_dirichlet")].iloc[0]
    assert nurse["P@5"] == pytest.approx(1 / 5) and nurse["MRR"] == 1.0 and nurse["nDCG@10"] == pytest.approx(1.0)       # exactly one nurse in the collection, ranked first
    S = R.summarize(df, ["lm_dirichlet", "hybrid"])
    assert (S["P@10"].between(0, 1)).all() and (S["nDCG@10_lo"] <= S["nDCG@10"]).all() and (S["nDCG@10"] <= S["nDCG@10_hi"]).all()
    cmp = R.compare(df, "hybrid", "hybrid")
    assert all(v["mean_diff"] == 0 and v["ties"] == 2 for v in cmp.values())
    md = render(df, QUERIES, st, "unit-test labels", 1, 4, "t", eng.index.n_docs)
    assert "no reliable difference" in md and "Pool depth: 4" in md and "ONE labeller" in md


def test_results_are_stored_in_ir_experiments(engine):
    from sqlalchemy import text
    df = pd.DataFrame([dict(system="hybrid", query_id="q1", query="python sql", latency_ms=3.0, **{"P@5": .4, "P@10": .3, "R@5": .5, "R@10": float("nan"), "MRR": 1.0, "nDCG@10": .7})])
    assert R.store_results(engine, df, "run-1", "manual", 100, {"thr": 1}) == 1
    assert R.store_results(engine, df, "run-1", "manual", 100, {"thr": 1}) == 1                       # re-running replaces, never duplicates
    with engine.connect() as c:
        row = c.execute(text("SELECT ranking_model, precision_at_k, recall_at_k, mrr, relevance_source FROM ir_experiments")).one()
        assert c.execute(text("SELECT COUNT(*) FROM ir_experiments")).scalar_one() == 1
    assert row == ("hybrid", 0.3, None, 1.0, "manual")


def test_report_refuses_conclusions_for_systems_with_unjudged_results():
    def rows(system, ndcg, judged):
        return [dict(system=system, query_id=f"q{i}", query=f"query {i}", n_results=10, latency_ms=2.0, **{"judged@10": judged, "P@5": .4, "P@10": .3, "R@5": .5, "R@10": .6,
                     "MRR": .5, "nDCG@10": ndcg + 0.01 * i, "n_relevant_judged": 4}) for i in range(8)]
    df = pd.DataFrame(rows("hybrid", 0.9, 1.0) + rows("lm_dirichlet", 0.2, 0.5) + rows("keyword_baseline", 0.3, 1.0))
    md = render(df, [{"id": f"q{i}", "text": f"query {i}", "need": ""} for i in range(8)], dict(judged=100, unlabeled=0, grade_counts={0: 50, 1: 30, 2: 20}), "unit-test", 1, 10, "t", 1000)
    assert "Unreliable systems" in md and "lm_dirichlet (only 50%" in md
    assert md.count("UNRELIABLE - a compared system has unjudged results") >= 3 and "supported after correcting" in md       # hybrid vs baseline is still evaluated


def test_labels_are_portable_across_databases(tmp_path):
    from jobintel.ir.index import IndexDoc, InvertedIndex
    eng, sy = systems_and_engine()
    pool = J.build_pool(sy, QUERIES, ["lm_dirichlet", "dfr_pl2"], 4)
    J.write_sheet(J.sheet_rows(pool, eng.index, eng.docstore), QUERIES, tmp_path / "s.xlsx")
    fill(tmp_path / "s.xlsx", lambda q, j: 1)
    labels, _ = J.read_labels(tmp_path / "s.xlsx")
    assert all(isinstance(k, tuple) and k[0] == "toy" and k[1].startswith("toy-") for d in labels.values() for k in d)
    # "another database": the same jobs but different internal job ids (ingestion order differs between machines)
    ix = eng.index
    shifted = [IndexDoc(9000 + i, [ix.terms[t] for t in ix.doc_tokens(i)],
                        {f: (None if isinstance(ix.meta[f][i], float) and ix.meta[f][i] != ix.meta[f][i] else ix.meta[f][i]) for f in ix.meta}) for i in range(ix.n_docs)]
    other = InvertedIndex.build(shifted)
    mapped, missing = J.resolve_labels(labels, other)
    assert missing == 0 and all(j >= 9000 for d in mapped.values() for j in d)
    assert sum(len(d) for d in mapped.values()) == sum(len(d) for d in labels.values())
    gone, missing = J.resolve_labels({"q1": {("toy", "toy-999"): 2}}, other)                     # a labeled job that this database does not have is reported, not guessed
    assert missing == 1 and gone == {}
