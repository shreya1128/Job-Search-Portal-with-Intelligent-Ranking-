"""The upload guard: every protection of the resume reader, tested against BOTH readers (the Phase 6b reader and the layout reader kept from the earlier snapshot)."""
import importlib
import io
import os
import threading
import zipfile
import shutil
from pathlib import Path

import pytest

from jobintel.rec import upload_guard as G
from jobintel.rec.upload_guard import UploadRejected, display_name, inspect, read_resume

READERS = ["jobintel.rec.documents"]
try:
    importlib.import_module("jobintel.rec.documents_layout")
    READERS.append("jobintel.rec.documents_layout")
except Exception:
    pass

SAMPLE_SCAN = Path(__file__).resolve().parents[1] / "data" / "samples" / "sample_resume_scan.png"


@pytest.fixture(params=READERS)
def reader(request):
    return importlib.import_module(request.param)


def text_pdf(pages: int, encrypt: bool = False) -> bytes:
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for p in range(pages):
        for i in range(30):
            c.drawString(60, 800 - 20 * i, f"Page {p + 1} line {i + 1}: Python SQL machine learning analyst Pune project delivered reports dashboards team")
        c.showPage()
    c.save()
    data = buf.getvalue()
    if encrypt:
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter()
        [w.add_page(p) for p in PdfReader(io.BytesIO(data)).pages]
        w.encrypt("secret")
        out = io.BytesIO()
        w.write(out)
        return out.getvalue()
    return data


def table_docx() -> bytes:
    import docx
    d = docx.Document()
    d.add_paragraph("Priya Sharma")
    d.add_paragraph("Skills: Python, SQL, Tableau")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Experience", "Data Analyst, Acme Pvt Ltd, 2019-2023"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def rewrite_docx(data: bytes, new_document_xml=None, extra_member=None) -> bytes:
    src, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for n in src.namelist():
            z.writestr(n, new_document_xml if (n == "word/document.xml" and new_document_xml is not None) else src.read(n))
        if extra_member:
            z.writestr(*extra_member)
    return out.getvalue()


def png(w: int, h: int) -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("L", (w, h), 255).save(buf, "PNG")
    return buf.getvalue()


# ------------------------------------------------------------------------------------------------------------ files that must be read
def test_good_files_are_read(reader):
    assert "Python" in read_resume(b"Skills: Python, SQL\nExperience: 3 years", "cv.txt", reader=reader).text
    d = read_resume(table_docx(), "cv.docx", reader=reader)
    assert "Data Analyst" in d.text and "Priya" in d.text
    p = read_resume(text_pdf(2), "cv.pdf", reader=reader)
    assert "Page 2 line 30" in p.text


@pytest.mark.skipif(not SAMPLE_SCAN.exists() or shutil.which("tesseract") is None, reason="sample scan or tesseract not present")
def test_scanned_image_goes_through_ocr_with_its_confidence(reader):
    d = read_resume(SAMPLE_SCAN.read_bytes(), "scan.png", reader=reader)
    assert len(d.text) > 300
    conf = getattr(d, "ocr_confidence", None) or getattr(d, "ocr_mean_conf", None)
    assert conf is not None and conf > 50


# ------------------------------------------------------------------------------------------------------------ refused, with a message a visitor can read
@pytest.mark.parametrize("name,data,fragment", [
    ("cv.txt", b"", "empty"),
    ("cv.exe", b"MZ" + os.urandom(50), "not supported"),
    ("cv.doc", b"\xd0\xcf\x11\xe0" + os.urandom(100), "Convert .doc"),
    ("noext", b"hello", "not supported"),
    ("cv.pdf", b"this is plain text, not a pdf\n" * 20, "not a PDF"),
    ("cv.pdf", b"MZ\x90\x00" + os.urandom(2000), "not a PDF"),
    ("cv.docx", b"plain text pretending to be a word file\n" * 20, "not a Word"),
    ("cv.png", b"plain text pretending to be an image\n" * 20, "not that kind of image"),
    ("cv.txt", os.urandom(5_000), "plain text"),
    ("cv.txt", b"Python\x00\x00\x00SQL" * 50, "plain text"),
])
def test_refused_files(reader, name, data, fragment):
    with pytest.raises(UploadRejected, match=fragment):
        read_resume(data, name, reader=reader)


def test_oversize_file_is_refused_before_any_parser_runs():
    class Boom:
        @staticmethod
        def read_document(_):
            raise AssertionError("the parser must not run for an oversize file")
    with pytest.raises(UploadRejected, match="larger than 15 MB"):
        read_resume(b"a" * (G.MAX_BYTES + 1), "cv.txt", reader=Boom)


def test_pdf_damaged_truncated_and_encrypted(reader):
    good = text_pdf(2)
    with pytest.raises(UploadRejected, match="damaged"):
        read_resume(good[:900], "cv.pdf", reader=reader)
    with pytest.raises(UploadRejected, match="password"):
        read_resume(text_pdf(2, encrypt=True), "cv.pdf", reader=reader)


def test_pdf_page_limit_passes_only_the_first_pages_on(reader, monkeypatch):
    monkeypatch.setattr(G, "MAX_PAGES", 3)
    d = read_resume(text_pdf(12), "cv.pdf", reader=reader)
    assert "Page 3 line 1:" in d.text and "Page 4 line 1:" not in d.text
    assert any("only the first 3 were read" in w for w in d.warnings)


def test_default_page_limit_is_eight_like_the_original_reader(reader):
    d = read_resume(text_pdf(20), "cv.pdf", reader=reader)
    assert "Page 8 line 1:" in d.text and "Page 9 line 1:" not in d.text


def test_docx_zip_bomb_by_total_size(reader, monkeypatch):
    monkeypatch.setattr(G, "MAX_DOCX_UNCOMPRESSED", 2_000_000)
    head = b'<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
    bomb = head + (b"<w:p><w:r><w:t>" + b"a" * 2000 + b"</w:t></w:r></w:p>") * 2500 + b"</w:body></w:document>"
    data = rewrite_docx(table_docx(), bomb)
    assert len(data) < 100_000 and len(bomb) > 4_000_000
    with pytest.raises(UploadRejected, match="expands"):
        read_resume(data, "cv.docx", reader=reader)


def test_docx_with_unsafe_member_path_or_missing_parts(reader):
    with pytest.raises(UploadRejected, match="unsafe"):
        read_resume(rewrite_docx(table_docx(), extra_member=("../../evil.txt", b"x")), "cv.docx", reader=reader)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "just a zip")
    with pytest.raises(UploadRejected, match="not a Word"):
        read_resume(buf.getvalue(), "cv.docx", reader=reader)
    with pytest.raises(UploadRejected, match="damaged"):
        read_resume(b"PK\x03\x04" + os.urandom(500), "cv.docx", reader=reader)


def test_image_pixel_limit_minimum_size_and_damage(reader, monkeypatch):
    monkeypatch.setattr(G, "MAX_PIXELS", 1_000_000)
    with pytest.raises(UploadRejected, match="too large"):
        read_resume(png(2000, 2000), "cv.png", reader=reader)
    with pytest.raises(UploadRejected, match="too small"):
        read_resume(png(80, 80), "cv.png", reader=reader)
    with pytest.raises(UploadRejected, match="could not be read"):
        read_resume(png(600, 600)[:200], "cv.png", reader=reader)


def test_the_9000_pixel_square_png_is_refused_with_the_real_limit():
    with pytest.raises(UploadRejected, match="too large"):
        inspect(png(9000, 9000), "bomb.png")


# ------------------------------------------------------------------------------------------------------------ what the parser and the server see
def test_parser_gets_a_private_constant_name_and_the_temp_file_is_removed():
    seen = {}

    class Spy:
        @staticmethod
        def read_document(path):
            seen["path"], seen["bytes"] = Path(path), Path(path).read_bytes()
            from types import SimpleNamespace
            return SimpleNamespace(text="Python", warnings=[])
    read_resume(b"Python", "../../C:\\Users\\victim\\cv<script>.txt", reader=Spy)
    assert seen["path"].name == "upload.txt" and seen["bytes"] == b"Python"
    assert not seen["path"].exists() and not seen["path"].parent.exists()


def test_parser_failures_become_generic_messages_without_paths():
    class Broken:
        @staticmethod
        def read_document(path):
            raise ValueError(f"cannot read {path} at /srv/jobintel/secret")
    with pytest.raises(UploadRejected) as e:
        read_resume(b"Python", "cv.txt", reader=Broken)
    assert "/srv" not in str(e.value) and "upload.txt" not in str(e.value) and "could not be read" in str(e.value)


def test_reader_messages_marked_safe_are_passed_on():
    class Safe:
        @staticmethod
        def read_document(path):
            raise type("DocumentError", (ValueError,), {})("the PDF has no pages")
    with pytest.raises(UploadRejected, match="no pages"):
        read_resume(b"Python", "cv.txt", reader=Safe)


def test_display_name_is_sanitised():
    assert display_name("..\\..\\evil<script>.pdf") == "evil_script_.pdf"
    assert display_name("/etc/passwd") == "passwd" and display_name("") == "resume" and len(display_name("a" * 300 + ".pdf")) == 80


def test_extracted_text_is_capped(monkeypatch):
    monkeypatch.setattr(G, "MAX_TEXT_CHARS", 50)

    class Big:
        @staticmethod
        def read_document(path):
            from types import SimpleNamespace
            return SimpleNamespace(text="x" * 500, warnings=[])
    d = read_resume(b"Python", "cv.txt", reader=Big)
    assert len(d.text) == 50 and any("cut to" in w for w in d.warnings)


def test_busy_server_says_so_instead_of_queueing_forever(monkeypatch):
    gate = threading.BoundedSemaphore(1)
    gate.acquire()
    monkeypatch.setattr(G, "_gate", gate)
    monkeypatch.setattr(G, "BUSY_WAIT_S", 0.05)
    with pytest.raises(UploadRejected, match="busy"):
        read_resume(b"Python", "cv.txt", reader=type("R", (), {"read_document": staticmethod(lambda p: None)}))
