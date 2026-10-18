"""Pooled relevance judgments.

POOLING: for every query the top-N results of every pooled system are merged into one list, SHUFFLED, and shown WITHOUT any information about
which system returned what (so the labeller cannot be biased by a system's rank).  The labeller marks each job 0 / 1 / 2.  Documents that no pooled
system returned are never judged and count as not relevant - the standard pooling assumption, with its known bias against systems that were
not pooled.  `judged@k` in the report shows how large that effect is for every evaluated system."""
from __future__ import annotations

import csv
import json
import random
from pathlib import Path
from typing import Callable, Dict, Iterable, List, Optional, Sequence, Tuple

GRADES = {0: "not relevant", 1: "partially relevant", 2: "clearly relevant"}
INSTRUCTIONS = [
    "HOW TO LABEL  (about 1-2 hours for 20 queries)",
    "",
    "For every row, read the query and its INFORMATION NEED (columns B-C) and decide how well THIS job posting answers it.",
    "Fill the 'relevance' column with one of:",
    "   2 = clearly relevant      the role and the key requirements in the query match what the posting is about",
    "   1 = partially relevant    related role or skills, but a key part of the need is missing (wrong seniority, wrong location when one was asked for, only one of the required skills ...)",
    "   0 = not relevant          a different kind of job, or the query words appear only incidentally",
    "",
    "Rules that keep the labels consistent:",
    " * Judge ONLY from what is shown in the row. Do not look anything up.",
    " * The rows of a query are shuffled and carry no information about which search method found them. Do not try to guess it.",
    " * If the query names a city or an experience level, a posting that does not satisfy it is at most 1.",
    " * Skill-only postings (key_skills filled, short snippet) are judged from the title and skills; postings whose text looks lower-cased without small words come from a source that stores text that way - judge the content, not the formatting.",
    " * Duplicated-looking postings are labelled independently.",
    " * Leave 'notes' empty unless something is unclear. Do not delete or reorder rows. Edit ONLY the 'relevance' (and 'notes') columns - the job id / source / external_id columns identify the job.",
    " * Do not leave any relevance cell empty: unlabeled rows make the evaluation refuse to run (or, with --allow-partial, bias it).",
    "",
    "When finished: save the file (keep it .xlsx) and send it back.",
]
HEADER = ["query_id", "query", "information_need", "job_id", "source", "external_id", "title", "company", "location", "experience", "key_skills", "description_snippet", "relevance", "notes"]


# ------------------------------------------------------------------------------------------------------------------ pool
def build_pool(systems: Dict[str, Callable], queries: Sequence[Dict], pool_systems: Sequence[str], depth: int, seed: int = 42) -> Dict:
    pool: Dict = {}
    for q in queries:
        per = {}
        for name in pool_systems:
            ids, _ = systems[name](q["text"])
            per[name] = [int(i) for i in ids[:depth]]
        union = list(dict.fromkeys(i for name in pool_systems for i in per[name]))
        random.Random(f"{seed}-{q['id']}").shuffle(union)
        pool[q["id"]] = {"query": q["text"], "need": q["need"], "systems": per, "job_ids": union}
    return pool


def save_pool(pool: Dict, path: Path) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(pool, indent=1), encoding="utf-8")


def load_pool(path: Path) -> Dict:
    return json.loads(Path(path).read_text(encoding="utf-8"))


# ------------------------------------------------------------------------------------------------------------------ sheet
def _doc_row(index, texts: Dict[int, Tuple[str, str]], job_id: int) -> Dict:
    d = index.doc_index(job_id)
    m = index.meta
    exp = None
    lo, hi = m["exp_min"][d], m["exp_max"][d]
    if m["source"][d] in ("naukri", "mcf", "toy") and lo == lo:              # experience STATED by the source
        exp = f"{lo:g}+ yrs" if hi != hi else (f"{lo:g} yrs" if lo == hi else f"{lo:g}-{hi:g} yrs")
    elif m["seniority"][d]:                                                    # other sources: only a seniority level (any years would be an approximation of ours)
        exp = f"level: {m['seniority'][d]}"
    title, text = texts.get(job_id, (m["title"][d], ""))
    q = m["text_quality"][d]
    snippet = (text or "").strip().replace("\n", " ")
    snippet = snippet[:700] + ("..." if len(snippet) > 700 else "")
    if q == "preprocessed":
        snippet = "[source stores lower-cased text without small words] " + snippet
    return dict(job_id=job_id, external_id=str(m["external_id"][d]) if m["external_id"][d] is not None else "", title=m["title"][d], company=m["company"][d] or "", location=", ".join(m["cities"][d] or ([m["city"][d]] if m["city"][d] else [])) or (m["country"][d] or ""),
                source=m["source"][d], experience=exp or "", key_skills=", ".join((m["skills"][d] or [])[:10]), description_snippet=snippet)


def sheet_rows(pool: Dict, index, docstore, only: Optional[Dict[str, set]] = None) -> List[Dict]:
    rows = []
    for qid, p in pool.items():
        ids = [j for j in p["job_ids"] if only is None or j in only.get(qid, set())]
        texts = docstore.get(ids) if ids else {}
        for j in ids:
            rows.append(dict(query_id=qid, query=p["query"], information_need=p["need"], relevance="", notes="", **_doc_row(index, texts, j)))
    return rows


def write_sheet(rows: List[Dict], queries: Sequence[Dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS, 1):
        ws.cell(i, 1, line)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 150
    wq = wb.create_sheet("Queries")
    wq.append(["query_id", "query", "information_need"])
    for q in queries:
        wq.append([q["id"], q["text"], q["need"]])
    for col, w in zip("ABC", (10, 50, 100)):
        wq.column_dimensions[col].width = w
    wj = wb.create_sheet("Judgments")
    wj.append(HEADER)
    for r in rows:
        wj.append([r.get(h, "") for h in HEADER])
    widths = dict(zip(HEADER, (9, 30, 38, 10, 10, 18, 34, 24, 22, 14, 40, 90, 12, 24)))
    for i, h in enumerate(HEADER, 1):
        wj.column_dimensions[wj.cell(1, i).column_letter].width = widths[h]
        wj.cell(1, i).font = Font(bold=True)
        wj.cell(1, i).fill = PatternFill("solid", fgColor="DDEBF7")
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in wj.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap
    rel_col = HEADER.index("relevance") + 1
    for r in range(2, wj.max_row + 1):
        wj.cell(r, rel_col).fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True, showErrorMessage=True, errorTitle="Invalid", error="Use 0, 1 or 2")
    wj.add_data_validation(dv)
    dv.add(f"{wj.cell(2, rel_col).coordinate}:{wj.cell(max(2, wj.max_row), rel_col).coordinate}")
    wj.freeze_panes = "E2"
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_csv(rows: List[Dict], path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=HEADER, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


# ------------------------------------------------------------------------------------------------------------------ read labels back
def _norm_ext(v) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_labels(path: Path) -> Tuple[Dict[str, Dict], Dict]:
    """-> ({query_id: {key: grade}}, stats).  key = (source, external_id) when the sheet has those columns (portable across databases built
    from the same raw files), otherwise the numeric job_id.  Raises ValueError listing the offending rows when a relevance cell is not 0/1/2."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Judgments"] if "Judgments" in wb.sheetnames else wb.worksheets[-1]
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding="utf-8-sig") as fh:
            grid = [list(r) for r in csv.reader(fh)]
    head = [str(h).strip().lower() if h is not None else "" for h in grid[0]]
    qi, ji, ri = head.index("query_id"), head.index("job_id"), head.index("relevance")
    si = head.index("source") if "source" in head else None
    ei = head.index("external_id") if "external_id" in head else None
    labels: Dict[str, Dict] = {}
    unlabeled: Dict[str, int] = {}
    bad: List[str] = []
    for n, row in enumerate(grid[1:], 2):
        if not row or row[qi] in (None, ""):
            continue
        qid, val = str(row[qi]).strip(), row[ri]
        key = (str(row[si]).strip(), _norm_ext(row[ei])) if si is not None and ei is not None and row[ei] not in (None, "") else int(float(row[ji]))
        if val in (None, ""):
            unlabeled[qid] = unlabeled.get(qid, 0) + 1
            continue
        try:
            g = int(float(str(val).strip()))
        except ValueError:
            g = -1
        if g not in GRADES:
            bad.append(f"row {n}: {val!r}")
            continue
        labels.setdefault(qid, {})[key] = g
    if bad:
        raise ValueError("relevance must be 0, 1 or 2 - problems: " + "; ".join(bad[:15]) + (" ..." if len(bad) > 15 else ""))
    counts = {g: sum(1 for q in labels.values() for v in q.values() if v == g) for g in GRADES}
    return labels, dict(judged=sum(counts.values()), unlabeled=sum(unlabeled.values()), unlabeled_by_query=unlabeled, grade_counts=counts, queries=len(labels))


def resolve_labels(labels: Dict[str, Dict], index) -> Tuple[Dict[str, Dict[int, int]], int]:
    """Map portable keys to THIS index's job ids.  Returns (labels by job id, number of labeled jobs that do not exist in this index)."""
    out: Dict[str, Dict[int, int]] = {}
    missing = 0
    for qid, d in labels.items():
        for key, g in d.items():
            di = index.doc_by_external(*key) if isinstance(key, tuple) else index.doc_index(key)
            if di is None:
                missing += 1
                continue
            out.setdefault(qid, {})[int(index.job_ids[di])] = g
    return out, missing


def keys_of(index, job_ids: Sequence[int]) -> set:
    return {(index.meta["source"][index.doc_index(j)], str(index.meta["external_id"][index.doc_index(j)])) for j in job_ids}
